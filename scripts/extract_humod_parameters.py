import openpyxl
import json

# Configure input and output files
input_path = "C:/Users/gh3270/Nextcloud/Working Folder/Tools/HU-MOD/humod2_expert_25-11-24_en_example.xlsx"
output_path_crops = "data/humod/crop_parameters.json"
output_path_fertiliser = "data/humod/fertiliser_parameters.json"

#Read excel file
humod_file = openpyxl.load_workbook(input_path, data_only= True)
print("Sheets:", humod_file.sheetnames)
crop_sheet = humod_file["HU-MOD-2-Parameters_CROP"] #Reading crop sheet
fertiliser_sheet = humod_file["HU-MOD-2-Parameters_FTLZ"] #Reading fertiliser sheet

# Extract crops
crops = {}

for row in crop_sheet.iter_rows(min_row=5, max_row=crop_sheet.max_row):

    crop_id = row[0].value
    name_eng = row[3].value
    name_de = row[2].value

    name = name_eng if name_eng else name_de

    if not crop_id or crop_id <= 0:
        continue
    if not name_eng and not name_de:
        continue
    if name in ("NN", "Dummy"):
        continue

    crops[str(crop_id)] = {
        "crop_id": int(crop_id),
        "name": name,
        "name_de": name_de or "",

        #Main product
        "main_product": {
            "dm_content_mp": row[4].value, #Dry matter content of main product
            "c_content_mp": row[5].value,  #C content of main product
            "n_content_mp": row[6].value   #N content of main product
        },
        
        #Side product
        "side_product": {
            "ratio_sp_mp": row[7].value,   #Ratio of side product to main product
            "dm_content_sp": row[8].value, #Dry matter content of side product
            "c_content_sp": row[9].value,  #C content of side product
            "n_content_sp": row[10].value  #N content of side product
        },

        #Litter
        "litter": {
            "ratio_lt_mp": row[11].value,   #Ratio of litter to main product
            "dm_content_lt": row[12].value, #Dry matter content of litter
            "c_content_lt": row[13].value,  #C content of litter
            "n_content_lt": row[14].value   #N content of litter
        },

        #Stubble
        "stubble": {
            "ratio_st_mp": row[15].value,  #Ratio of stubble to main product
            "dm_content_st": row[16].value, #Dry matter content of stubble
            "c_content_st": row[17].value,  #C content of stubble
            "n_content_st": row[18].value  #N content of stubble           
        },

        #Root
        "root": {
            "ratio_root_shoot": row[19].value,  #Ratio of root:shoot
            "dm_content_rt": row[20].value, #Dry matter content of root
            "c_content_rt": row[21].value,  #C content of root
            "n_content_rt": row[22].value  #N content of root           
        },

        #Legume N fixation
        "legume_fixation": {
            "leg_share": row[23].value, #Legume share in case of mixed crops or pure crops (0 or 1)
            "NDFA": row[24].value       #N derived from atmosphere
        },

        #N dynamics
        "n_dynamics": {
            "NYRN": row[25].value,      #Non-yield related N lost in cultivation
            "NAV": row[26].value        #N availability factor
        },



        ## VDLUFA humus balance coefficients
        "vdlufa": {
            "humus_supply_low": row[27].value, # Humus supply coefficient - lower limit
            "humus_supply_high": row[28].value, # Humus supply coefficient - higher limit
            "humus_supply_organic": row[29].value, # Humus supply coefficient - for organic farms
            "supply_coeff_sp": row[31].value, # Humus supply coefficient for the crop's side product (from the original VDLUFA parameter)
        },


        ## Reference yield
        "default_yield": row[30].value/10 if row[30].value else None # Default yield of selected crop, converted to tonnes (from HU-MOD)
    }


# Extract fertilisers
fertilisers = {}

for row in fertiliser_sheet.iter_rows(min_row=5, max_row=fertiliser_sheet.max_row):
    fert_id = row[0].value
    name_fert_eng = row[3].value
    name_fert_de = row[2].value

    name_fert = name_fert_eng if name_fert_eng else name_fert_de

    if not fert_id or fert_id <= 0:
        continue
    if not name_fert_eng and not name_fert_de:
        continue
    if row[4].value is None and not name_fert:
        continue

    fertilisers[str(fert_id)] = {
        "fert_id": int(fert_id),
        "name": name_fert,

        "dm_content_fert": row[4].value, # Dry matter content of fertiliser
        "c_content_fert": row[5].value, # C content of fertiliser
        "n_content_fert": row[6].value, # N content of fertiliser
        "NFAV": row[7].value, # Fraction of N available to crop
        "NFAV_orig": row[17].value, # Fraction of N available to crop (before recalculation)

        ## VDLUFA humus balance coefficients
        "vdlufa_fert": {
            "supply_coeff_fert": row[8].value,
            "supply_coeff_sp_updated": row[9].value # Humus supply coefficient for the crop's side product (recalculated as organic input corresponding crop by HU-MOD)
        }

    }

#save crops to json
with open(output_path_crops, "w", encoding= "utf-8") as f:
    json.dump(crops, f, indent=2, ensure_ascii=False)
print(f"saved {len(crops)} crops")

#save fertilisers to json
with open(output_path_fertiliser, "w", encoding= "utf-8") as f:
    json.dump(fertilisers, f, indent=2, ensure_ascii=False)
print(f"saved {len(fertilisers)} fertilisers")




